#!/usr/bin/env python3
"""Generate resume.docx matching the website content. Fits on one page."""

import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor, Twips
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn

doc = Document()

# Tight page margins for one-page fit
for section in doc.sections:
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.4)
    section.left_margin = Inches(0.6)
    section.right_margin = Inches(0.6)

style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(9.5)
style.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
style.paragraph_format.space_after = Pt(0)
style.paragraph_format.space_before = Pt(0)
style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

NAVY = RGBColor(0x13, 0x27, 0x43)
TEAL = RGBColor(0x0D, 0x94, 0x88)
MUTED = RGBColor(0x47, 0x55, 0x69)
LIGHT = RGBColor(0x94, 0xA3, 0xB8)


def add_heading_styled(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = NAVY
    pPr = p._p.get_or_add_pPr()
    pBdr = pPr.makeelement(qn("w:pBdr"), {})
    bottom = pBdr.makeelement(qn("w:bottom"), {
        qn("w:val"): "single",
        qn("w:sz"): "6",
        qn("w:color"): "0D9488",
        qn("w:space"): "1",
    })
    pBdr.append(bottom)
    pPr.append(pBdr)
    return p


def add_bullet(text, size=9):
    p = doc.add_paragraph(text, style="List Bullet")
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    for run in p.runs:
        run.font.size = Pt(size)
    return p


# --- Name ---
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(1)
run = p.add_run("NICHOLAS BRADFORD")
run.bold = True
run.font.size = Pt(18)
run.font.color.rgb = NAVY

# --- Title ---
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(1)
run = p.add_run("AI & Analytics Professional | @XeoMatrix, Ex-Salesforce, Ex-Amazon")
run.font.size = Pt(10)
run.font.color.rgb = MUTED

# --- Contact ---
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(4)
run = p.add_run("678-646-2738  |  nicholasfbradford@gmail.com  |  LinkedIn")
run.font.size = Pt(9)
run.font.color.rgb = MUTED

# --- Profile ---
add_heading_styled("PROFILE")
profile_items = [
    'Exceeded FY 2024 sales quota by 120% at Salesforce/Tableau, winning "Rocky of the Quarter"',
    "Architected auto-scaling data pipelines and deep-learning models flagging product-safety defects with 98%+ precision, averting $45M+ in recall/litigation exposure",
    "Boosted incremental revenue $1.2M (+4%) and saved $5M labor cost by deploying prescriptive models at Comcast",
    "Architected AI-driven labor-optimization demand models; reduced overtime 18% and unlocked $6M in annual savings via Tableau & Power BI",
]
for item in profile_items:
    add_bullet(item)

# --- Skills ---
add_heading_styled("SKILLS")
skills = [
    "Python & R (Certified) \u2013 PyTorch, TensorFlow, XGBoost",
    "Tableau (Certified)  |  AWS, Azure, Hive, Databricks, Spark SQL, DBT, FiveTran, Snowflake (Certified)",
    "MS Office \u2013 VBA, Excel, Access, Power Tools  |  ETL & ELT  |  Salesforce",
]
for s in skills:
    add_bullet(s)

# --- Experience ---
add_heading_styled("EXPERIENCE")

jobs = [
    ("Lead Analytics Consultant", "XeoMatrix", "Sep 2025 \u2013 Present", [
        "Directed a high-performance team of 2 Senior Consultants, overseeing end-to-end delivery of $500k in cloud-native AI projects while maintaining a 100% client retention rate",
        "Architected AWS/GCP-native data foundations for enterprise clients, integrating Snowflake and Databricks to reduce reporting latency by 60% and cloud TCO by $1.2M annually",
        "Spearheaded $350k in net-new ARR by leading technical sales cycles, delivering AI-driven proof-of-concepts integrating LLMs with Tableau Pulse for C-suite decision-makers",
        "Designed production-grade RAG pipelines on AWS SageMaker, surfacing insights via embedded Tableau dashboards to automate executive reporting",
    ]),
    ("Sr. Solutions Analytics Architect", "Salesforce / Tableau", "May 2022 \u2013 Sep 2025", [
        "Closed $10M+ in net-new ARR by architecting and live-demoing AI analytics solutions for mid-market & enterprise accounts, serving as trusted C-suite advisor",
        "Led AI enablement: prompt-engineering masterclass to 450 engineers; introduced AI and LLM tools adopted org-wide globally (2023\u201324)",
        "Built AI pipelines: GPT prompt design, vector search RAG, Snowflake ELT, real-time dashboards, full MLOps for embedded analytics",
        "Created workshops for customers on Tableau, Salesforce and effective data strategy implementation",
    ]),
    ("Data Engineer", "Amazon", "Sep 2020 \u2013 May 2022", [
        "Built data pipelines using Cradle, Datanet, AWS Lambda, and Redshift; end-to-end reporting and analytics solutions",
        "Built AI NLP models surfacing product-safety outliers; briefed S-Team (Amazon Senior Leadership) bi-weekly",
        "Designed high-performance pipeline improving data processing time by 40%; reduced storage costs by 60% and cut $70K in cloud spend",
    ]),
    ("Sr. Data Scientist (Promotion)", "Comcast / Universal Orlando", "Jul 2018 \u2013 Aug 2020", [
        "Architected AI-driven optimization pipelines (Spark & Databricks) with reinforcement learning and demand models; reduced overtime 18%, unlocked $6M in annual savings",
        "Deployed Spark-based demand models saving $5M annually; created incentive strategy increasing revenue by $1.2M (+~4%)",
    ]),
    ("Division Lead Sr. BI Analyst", "Comcast", "Jun 2015 \u2013 Jul 2018", [
        "Innovation Honor Roll Award for modernizing workflows and predictive analysis via SAP data warehouse (BW)",
        "Designed reports, visualizations and dashboards leveraging Tableau/Power BI, Excel, SAP BO, SQL Teradata and Python",
    ]),
    ("Logistics Manager", "Goodman Networks", "Jun 2014 \u2013 Jun 2015", [
        "Automated shipping ETA updates with Python, streamlining asset provisioning",
    ]),
    ("Associate Equity Research Analyst", "Helm Bank Equity", "Jun 2012 \u2013 Jun 2015", [
        "Equity-research analysis",
    ]),
]

for title, company, dates, bullets in jobs:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    run = p.add_run(title)
    run.bold = True
    run.font.size = Pt(9.5)
    run = p.add_run("  |  " + company)
    run.font.size = Pt(9.5)
    run.font.color.rgb = MUTED
    run = p.add_run("  |  " + dates)
    run.font.size = Pt(8.5)
    run.font.color.rgb = LIGHT

    for b in bullets:
        add_bullet(b)

# --- Education ---
add_heading_styled("EDUCATION")
for school, degree, date in [
    ("Georgia Institute of Technology", "Master of Science in Analytics", "August 2020"),
    ("Sergio Arboleda University", "B.S. Finance", "May 2013"),
]:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    run = p.add_run(school)
    run.bold = True
    run.font.size = Pt(9.5)
    run = p.add_run("  \u2014  " + degree)
    run.font.size = Pt(9.5)
    run.font.color.rgb = MUTED
    run = p.add_run("  (" + date + ")")
    run.font.size = Pt(8.5)
    run.font.color.rgb = LIGHT

# --- Certifications ---
add_heading_styled("CERTIFICATIONS")
certs = [
    "Snowflake \u2014 Essentials (Jan 2023)  |  Tableau \u2014 Data Scientist (May 2020)",
    "Enthought \u2014 Python for Data Analysis (Sep 2018)  |  Thinkful \u2014 Data Science in Python (Apr 2015)",
]
for c in certs:
    add_bullet(c)

# Save
out_dir = os.path.join(os.path.dirname(__file__), "..", "assets")
os.makedirs(out_dir, exist_ok=True)
out_path = os.path.join(out_dir, "resume.docx")
doc.save(out_path)
print(f"Generated: {out_path}")
